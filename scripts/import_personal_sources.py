"""Import food reference values and a frozen workout plan, never execute source formulas.

Usage: python import_personal_sources.py --workbook PLAN.xlsx --diaries MAY.pdf JUNE.pdf
Raw documents stay local. Product nutrition is read by column position, not token order.
"""
import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'app/src/main/assets'


def import_plan(path):
    formulas = openpyxl.load_workbook(path, data_only=False)['1 в день']
    values = openpyxl.load_workbook(path, data_only=True)['1 в день']
    bases = {4: 'squat', 5: 'bench', 22: 'deadlift'}

    def base_for(row, seen=None):
        if row in bases:
            return bases[row]
        seen = set() if seen is None else seen
        if row in seen:
            raise ValueError(f'Cyclic maximum reference at AH{row}')
        seen.add(row)
        expr = str(formulas.cell(row, 34).value)
        refs = re.findall(r'\$?AH\$?(\d+)', expr)
        if not refs:
            raise ValueError(f'Unsupported maximum formula AH{row}: {expr}')
        return base_for(int(refs[0]), seen)

    days = []
    checks = []
    for cycle in range(25):
        for day in range(6):
            start = 4 + cycle * 56 + day * 9
            exercises = []
            for row in range(start, min(start + 8, 1402)):
                name = values.cell(row, 4).value
                if not name or name == 'Упражнения':
                    continue
                groups = []
                for col in [7, 11, 15, 19, 23, 27]:
                    weight, reps, sets, percent = [values.cell(row, col+i).value for i in range(4)]
                    if all(isinstance(n, (int, float)) and n > 0 for n in (reps, sets, percent)):
                        groups.append({'reps': int(reps), 'sets': int(sets), 'percent': percent})
                        checks.append({'row': row, 'group': len(groups)-1, 'expectedWeight': weight})
                if not groups:
                    continue
                base = base_for(row)
                base_value = {'squat': 230, 'bench': 160, 'deadlift': 240}[base]
                multiplier = values.cell(row, 6).value
                exercises.append({
                    'id': row, 'name': name, 'load': values.cell(row, 2).value or '',
                    'base': base, 'factor': values.cell(row, 34).value / base_value,
                    'rounding': 1.0 if multiplier == 2 else (2.5 if multiplier in (None, '', 1) else 0.1),
                    'equipmentMultiplier': multiplier if isinstance(multiplier, (int, float)) else 1,
                    'groups': groups,
                })
            if exercises:
                comment = values.cell(start, 38).value
                days.append({'id': len(days), 'cycle': cycle+1, 'day': day+1,
                             'notes': comment if isinstance(comment, str) else '', 'exercises': exercises})
    payload = {'title': 'СРЦ7 · 1 в день', 'sourceUrl': 'https://docs.google.com/spreadsheets/d/1XUjSjMRlWZDY_bSzxbbpLvxjwk6fR6UKCpmhROFg8L8/edit#gid=848117894',
               'sourceDate': '2026-10-19', 'referenceMaxima': {'squat': 230, 'bench': 160, 'deadlift': 240}, 'days': days}
    (ASSETS / 'workout-plan.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    (ASSETS / 'workout-reference-checks.json').write_text(json.dumps(checks), encoding='utf-8')
    print(f'Workout: {len(days)} days, {sum(len(d["exercises"]) for d in days)} exercises, {len(checks)} weight checks', flush=True)


def import_diaries(paths):
    samples = defaultdict(list)
    for path in paths:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages[1:]:
                words = page.extract_words()
                # Narrow labels can touch the calorie cell: pdfplumber returns "...710".
                # Split this known clipping boundary and keep the printed product name.
                for word in list(words):
                    fused = re.fullmatch(r'(\.\.\.)(\d+(?:,\d+)?)', word['text'])
                    if fused and 200 < word['x0'] < 239:
                        word['text'] = fused[1]
                        words.append({**word, 'text': fused[2], 'x0': 239, 'x1': 255, 'top': word['top']+2})
                lines = defaultdict(list)
                for word in words:
                    lines[round(word['top']/4)*4].append(word)
                ordered = sorted(lines.items())
                for idx, (y, line) in enumerate(ordered):
                    nums = [w for w in line if w['x0'] > 230 and re.fullmatch(r'\d+(?:,\d+)?', w['text'])]
                    if not nums:
                        continue
                    # Product labels sit 1-2pt above their numeric row, quantities below it.
                    nearby = [w for w in words if 0.5 < nums[0]['top'] - w['top'] <= 3 and w['x0'] < 239]
                    name = ' '.join(w['text'] for w in sorted(nearby, key=lambda w: w['x0']))
                    if not name or name.startswith(('Всего', 'Страница')):
                        continue
                    nutrients = {}
                    for w in nums:
                        center = (w['x0'] + w['x1'])/2
                        col = round((center-247)/60)
                        nutrients[col] = float(w['text'].replace(',', '.'))
                    if 0 not in nutrients:
                        continue
                    amount_words = sorted([w for w in words if 8 < w['top']-nums[0]['top'] < 21 and w['x0'] < 235], key=lambda w: w['x0'])
                    portion = ' '.join(w['text'] for w in amount_words)
                    quantities = re.findall(r'(\d+(?:[,.]\d+)?)\s*(г|мл)\b', portion)
                    normalized_portion = re.sub(r'(?<=\d)\s+(?=\d{3}\b)', '', portion)
                    quantities = re.findall(r'(\d+(?:[,.]\d+)?)\s*(г|мл)\b', normalized_portion)
                    amount = float(quantities[-1][0].replace(',', '.')) if quantities else None
                    unit = quantities[-1][1] if quantities else ''
                    # FatSecret can omit blank nutrient columns. Never shift later values left.
                    if not all(col in nutrients for col in (0, 1, 3, 6)):
                        raise ValueError(f'Missing required nutrition columns: {name}: {nutrients}')
                    samples[name].append({'values': [nutrients[c] for c in (0, 6, 1, 3)],
                                          'amount': amount, 'unit': unit, 'portion': portion})
                page.flush_cache()
        print(f'Read {Path(path).name}: {len(samples)} distinct names so far', flush=True)
    products = []
    for name, rows in sorted(samples.items()):
        measured = [s for s in rows if s['amount'] and s['amount'] > 0]
        sample = measured[0] if measured else rows[0]
        divisor = sample['amount']/100 if measured else 1
        nutrition = [round(v/divisor, 3) for v in sample['values']]
        note = 'Из пищевых дневников за май и июнь 2026. Значения FatSecret.'
        if '...' in name:
            note += ' Название сокращено в исходном PDF.'
        if measured and sample['unit'] == 'мл':
            note += ' Пищевая ценность на 100 мл; вводите объём в мл.'
        products.append({'name': name, 'calories': nutrition[0], 'protein': nutrition[1], 'fat': nutrition[2], 'carbs': nutrition[3],
                         'portionLabel': sample['portion'] if not measured else '', 'unit': sample['unit'] or 'порц.',
                         'aliases': {'Dimljena Pecenica': 'Димлена печеница копченая свинина pečenica Сербия Zlatiborac',
                                     'Protein Imlek': 'Протеин Имлек белковый напиток Сербия',
                                     'Grafschafter High Protein Wraps': 'Графшафтер протеиновые тортильи лаваш Lidl Лидл',
                                     'High Protein Wraps': 'Протеиновые тортильи лаваш Lidl Лидл',
                                     'Milbona Моцарелла Light': 'Милбона моцарела mozzarella light Lidl Лидл Сербия'}.get(name, ''),
                         'source': 'diary', 'note': note})
    (ASSETS / 'diary-products.json').write_text(json.dumps(products, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Diary: {len(products)} products, {sum(bool(p["portionLabel"]) for p in products)} portion-based (no invented gram weights)', flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook')
    parser.add_argument('--diaries', nargs='+')
    args = parser.parse_args()
    ASSETS.mkdir(parents=True, exist_ok=True)
    if args.workbook:
        import_plan(args.workbook)
    if args.diaries:
        import_diaries(args.diaries)
